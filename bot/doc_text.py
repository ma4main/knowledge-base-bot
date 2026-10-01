"""Извлечение текста из присланных файлов: сырьё для базы знаний.

Разбираем только текстовые форматы и office-документы без картинок: .txt/.md/.csv
как есть, .docx — абзацы и таблицы (python-docx), .pdf — текстовый слой (pypdf),
без OCR. Файл в библиотеку кладётся отдельно от разбора.
"""

from __future__ import annotations

import logging
from pathlib import Path

log = logging.getLogger(__name__)

PLAIN_EXTS = {".txt", ".md", ".markdown", ".csv", ".log", ".srt", ".vtt"}
DOCX_EXTS = {".docx"}
PDF_EXTS = {".pdf"}
SUPPORTED = PLAIN_EXTS | DOCX_EXTS | PDF_EXTS

# Длинное режем: в модель текст уходит вместе с каталогом базы.
MAX_CHARS = 40000


def can_extract(ext: str) -> bool:
    return ext.lower() in SUPPORTED


def extract(path: Path, ext: str) -> tuple[str, str]:
    """Возвращает (текст, примечание для человека); пустой текст — разобрать не удалось."""
    ext = ext.lower()
    try:
        if ext in PLAIN_EXTS:
            text = _plain(path)
        elif ext in DOCX_EXTS:
            text = _docx(path)
        elif ext in PDF_EXTS:
            text = _pdf(path)
        else:
            return "", f"формат {ext} не разбираю"
    except Exception as error:
        # Битый формат — ожидаемая ситуация, трейсбэк в логах ни к чему.
        log.warning("Не смог разобрать %s (%s): %s", path.name, ext, error)
        return "", "файл не удалось разобрать"

    text = (text or "").strip()
    if not text:
        if ext in PDF_EXTS:
            return "", "в PDF нет текстового слоя (похоже, скан) — распознавание не делаю"
        return "", "текста в файле не нашёл"
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS].rsplit("\n", 1)[0]
        return text, f"файл длинный — взял первые {len(text)} символов"
    return text, ""


def _plain(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        # Выгрузки из Windows-программ часто в cp1251.
        return path.read_text(encoding="cp1251", errors="replace")


def _docx(path: Path) -> str:
    from docx import Document  # импорт внутри: нужен только для .docx

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pdf(path: Path) -> str:
    from pypdf import PdfReader  # импорт внутри: нужен только для .pdf

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)
